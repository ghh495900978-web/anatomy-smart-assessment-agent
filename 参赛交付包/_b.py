# -*- coding: utf-8 -*-
import io, re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
doc = Document()
st = doc.styles["Normal"]; st.font.name = "宋体"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
doc.styles["Normal"].paragraph_format.space_after = Pt(6); doc.styles["Normal"].paragraph_format.line_spacing = 1.5
for s in doc.sections:
    s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5); s.left_margin = Cm(3.0); s.right_margin = Cm(3.0)
def cjk(r, f="宋体"):
    r.font.name = f; r._element.rPr.rFonts.set(qn("w:eastAsia"), f)
lines = io.open("01_教学智能体建设说明书.md", encoding="utf-8").read().split("\n")
i = 0
while i < len(lines):
    s = lines[i].strip()
    if not s or (s.startswith("---") and set(s) <= set("-")): i += 1; continue
    if s.startswith("# "):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(s[2:].strip()); r.bold = True; r.font.size = Pt(18); cjk(r, "黑体"); i += 1; continue
    if s.startswith("## "):
        p = doc.add_paragraph(); r = p.add_run(s[3:].strip()); r.bold = True; r.font.size = Pt(14); cjk(r, "黑体"); p.paragraph_format.space_before = Pt(12); i += 1; continue
    if s.startswith("### "):
        p = doc.add_paragraph(); r = p.add_run(s[4:].strip()); r.bold = True; r.font.size = Pt(12.5); p.paragraph_format.space_before = Pt(8); i += 1; continue
    if s.startswith("|") and i+1 < len(lines) and re.match(r"^[\s:\-\|]+$", lines[i+1].strip()):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not re.match(r"^[\s:\-]+$", "".join(cells)): rows.append(cells)
            i += 1
        if rows:
            ncol = max(len(r) for r in rows); t = doc.add_table(rows=0, cols=ncol); t.style = "Table Grid"
            for ri, row in enumerate(rows):
                cs = t.add_row().cells
                for ci in range(ncol):
                    txt = re.sub(r"\*\*(.+?)\*\*", r"\1", row[ci] if ci < len(row) else "")
                    cs[ci].text = ""; run = cs[ci].paragraphs[0].add_run(txt); run.font.size = Pt(10.5)
                    if ri == 0: run.bold = True
        continue
    if s.startswith(">"):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8)
        r = p.add_run(s.lstrip("> ").strip()); r.font.size = Pt(11); r.font.color.rgb = RGBColor(0x44,0x44,0x44); i += 1; continue
    p = doc.add_paragraph(); p.paragraph_format.first_line_indent = Cm(0.74)
    txt = re.sub(r"\*\*(.+?)\*\*", r"\1", s); txt = re.sub(r"`(.+?)`", r"\1", txt)
    r = p.add_run(txt); r.font.size = Pt(12); i += 1
doc.save("01_教学智能体建设说明书.docx")
print("OK")
